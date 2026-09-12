"""Read-only validation for local debtor files. Never writes customer records."""
import csv
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
ALIASES = {
    "invoice": ("invoice no", "invoice number", "invoice"),
    "name": ("name", "customer", "customer name"),
    "email": ("email", "e mail", "email address", "contact email"),
    "purchase_date": ("invoice date", "purchase date"),
    "due_date": ("due date",),
    "invoice_amount": ("invoice amount", "original amount"),
    "amount_paid": ("amount paid", "paid"),
}

def normalized(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()

def iso_date(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try: return datetime.strptime(str(value).strip(), fmt).date().isoformat()
        except ValueError: pass
    raise ValueError("Invalid date. Use DD/MM/YYYY or YYYY-MM-DD.")

def amount(value):
    try:
        number = Decimal(str(value).replace("R", "").replace(",", "").replace(" ", ""))
        if not number.is_finite() or number < 0 or number > 999999999 or number != number.quantize(Decimal(".01")):
            raise ValueError()
        return str(number.quantize(Decimal(".01")))
    except (ValueError, InvalidOperation):
        raise ValueError("Invalid amount.")

def read_snapshot(path, sheet=None):
    """Return normalized rows, with explicit rejection of ambiguous/incomplete data."""
    path = Path(path)
    if path.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError("Select a CSV or XLSX file.")
    before = path.stat()
    if before.st_size > MAX_BYTES: raise ValueError("File exceeds 10 MB.")
    raw = path.read_bytes()
    after = path.stat()
    if (before.st_size,before.st_mtime_ns) != (after.st_size,after.st_mtime_ns):
        raise ValueError("File is still being saved; retry on the next check.")
    workbook = None
    if path.suffix.lower() == ".csv":
        text = raw.decode("utf-8-sig")
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error: dialect = csv.excel
        records = iter(csv.reader(io.StringIO(text), dialect))
    else:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 50 MB.")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        if not sheet and len(workbook.sheetnames) != 1:
            workbook.close()
            raise ValueError("Choose a worksheet by name.")
        try: worksheet = workbook[sheet] if sheet else workbook.active
        except KeyError:
            workbook.close()
            raise ValueError("Worksheet not found.")
        records = worksheet.iter_rows(values_only=True)
    try:
        headers = next(records, ())
        mapping = {}
        for field, names in ALIASES.items():
            matches = [i for i,h in enumerate(headers) if normalized(h) in names]
            if len(matches)>1: raise ValueError(f"Ambiguous columns for {field}.")
            if matches: mapping[field] = matches[0]
        if not {"invoice", "name"} <= mapping.keys():
            raise ValueError("Invoice No. and Name columns are required.")
        rows, seen = [], set()
        for number, record in enumerate(records, 2):
            if number > MAX_ROWS+1: raise ValueError("File exceeds 10,000 rows.")
            if not any(v is not None and str(v).strip() for v in record): continue
            row = {k: record[i] for k,i in mapping.items() if i<len(record) and record[i] is not None and str(record[i]).strip()}
            if not row.get("invoice") or not row.get("name"):
                raise ValueError(f"Row {number}: invoice and customer name are required.")
            row["invoice"] = str(row["invoice"]).strip()
            row["name"] = str(row["name"]).strip()
            if row["invoice"] in seen: raise ValueError(f"Row {number}: duplicate invoice.")
            seen.add(row["invoice"])
            if "email" in row:
                row["email"] = str(row["email"]).strip()
                if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+",row["email"]):
                    raise ValueError(f"Row {number}: invalid email.")
            for key in ("purchase_date", "due_date"):
                if key in row: row[key] = iso_date(row[key])
            for key in ("invoice_amount", "amount_paid"):
                if key in row: row[key] = amount(row[key])
            if "invoice_amount" in row and "amount_paid" in row:
                balance = Decimal(row["invoice_amount"]) - Decimal(row["amount_paid"])
                if balance<0: raise ValueError(f"Row {number}: amount paid exceeds invoice amount.")
                row["balance"] = str(balance)
            rows.append(row)
        return rows
    finally:
        if workbook: workbook.close()
